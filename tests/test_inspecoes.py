import contextlib
import hashlib
import io
import shutil
import sys
import uuid
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
import etl_inspecoes as etl


class Inspecoes(unittest.TestCase):
    def setUp(self):
        self.root = etl.ROOT / 'work_tests' / uuid.uuid4().hex
        self.root.mkdir(parents=True)
        self.addCleanup(shutil.rmtree, self.root)
        self.entry = self.root / 'entrada'
        shutil.copytree(etl.ROOT / 'dados/entrada_inspecoes', self.entry)
        self.out = self.root / 'saida'

    def read(self):
        return etl.ler_validar(self.entry, self.out)

    def mutate(self, cell, value):
        p = self.entry / 'inspecao_sintetica_1.xlsx'
        w = openpyxl.load_workbook(p)
        w.active[cell] = value
        w.save(p)
        w.close()

    def test_export_and_hashes(self):
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(etl.run(self.entry, self.out), 0)
        w = openpyxl.load_workbook(self.out / 'Consolidado.xlsx')
        self.addCleanup(w.close)
        self.assertEqual(w.sheetnames, ['Consolidado', 'Validacao', 'Origens'])
        self.assertEqual(w['Consolidado'].max_row, 41)
        self.assertEqual(w['Origens'].max_row, 3)
        self.assertIn(0, [r[8] for r in list(w['Consolidado'].values)[1:]])
        for r in list(w['Origens'].values)[1:]:
            self.assertEqual(r[2], hashlib.sha256((self.entry/r[0]).read_bytes()).hexdigest())

    def test_invalid_notes_preserve_previous(self):
        self.out.mkdir()
        target = self.out / 'Consolidado.xlsx'
        for note in (None, -1, 6, 1.5, '=1+1'):
            with self.subTest(note=note):
                target.write_bytes(b'previous')
                self.mutate('D5', note)
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(etl.run(self.entry, self.out), 2)
                self.assertEqual(target.read_bytes(), b'previous')

    def test_copy_and_stable_ids(self):
        first = self.read()[0]['dInspecoes']
        p = self.entry / 'inspecao_sintetica_1.xlsx'
        shutil.copy(p, self.entry/'copia.xlsx')
        tables, issues, _, summary = self.read()
        self.assertEqual(summary['erros'], 0)
        self.assertEqual(first, tables['dInspecoes'])
        self.assertTrue(any(i['CODIGO']=='COPIA_IGNORADA' for i in issues))
        p.rename(self.entry/'renomeado.xlsx')
        self.assertEqual(first, self.read()[0]['dInspecoes'])

    def test_conflicting_copy(self):
        shutil.copy(self.entry/'inspecao_sintetica_1.xlsx', self.entry/'copia.xlsx')
        self.mutate('D5', 5)
        self.assertTrue(any(i['CODIGO']=='INSPECAO_CONFLITANTE' for i in self.read()[1]))

    def test_structure_and_metadata(self):
        for cell, value, code in [('E2','DESCONHECIDA','SIGLA_DESCONHECIDA'), ('B3','31/02/2026','DATA_INVALIDA'), ('G5','alterado','CATALOGO_DIVERGENTE'), ('A1','outro','MODELO_NAO_APROVADO')]:
            with self.subTest(code=code):
                p = self.entry/'inspecao_sintetica_1.xlsx'
                shutil.copy(etl.ROOT/'dados/entrada_inspecoes'/p.name, p)
                self.mutate(cell, value)
                self.assertTrue(any(i['CODIGO']==code for i in self.read()[1]))

    def test_empty_and_output_guard(self):
        for p in self.entry.iterdir():
            p.unlink()
        self.assertGreater(self.read()[3]['erros'], 0)
        with self.assertRaises(ValueError):
            etl.ler_validar(self.entry, self.entry/'saida')

    def test_formula_injection_is_text(self):
        self.mutate('E3', '=HYPERLINK("https://example.invalid")')
        self.assertGreater(self.read()[3]['erros'], 0)
        self.mutate('E3', '@avaliador ficticio')
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(etl.run(self.entry, self.out), 0)
        w = openpyxl.load_workbook(self.out/'Consolidado.xlsx')
        self.addCleanup(w.close)
        cells = [c for row in w['Consolidado'] for c in row if c.value=='@avaliador ficticio']
        self.assertTrue(cells)
        self.assertTrue(all(c.data_type=='s' for c in cells))
